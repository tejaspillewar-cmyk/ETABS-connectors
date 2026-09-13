# %%
# =====================================================================
#  FDR AUTOMATION TOOL — Integrated with Etabs_live.py
#  Python port of FDR-tool-t2 (React/TypeScript)
#
#  This module connects to ETABS via Etabs_live.py, extracts pier
#  force data & geometry, and performs:
#    1. Required Pt% (steel percentage) calculation
#    2. 0.4 fck Capacity/Demand check
#    3. 0.2 fck Boundary (ductile detailing) check
#    4. CAD export script generation (AutoCAD PLINE + TEXT commands)
#    5. Excel report export
#
#  Usage:
#      # As an importable module:
#      from Etabs_live import connect_etabs
#      from fdr_tool import FDRTool, FDRConfig
#      _, SapModel = connect_etabs()
#      tool = FDRTool(SapModel)
#      results = tool.run_all()
#
#      # Or run directly:
#      python fdr_tool.py
#
#      # Or run cells interactively in VS Code (Ctrl+Enter on # % blocks)
# =====================================================================

from __future__ import annotations

import math
import os
import re
from dataclasses import dataclass, field
from datetime import datetime
from typing import Optional

import pandas as pd


# %%
# =================================================================
# FDRConfig — All engineering constants in one place
# =================================================================

@dataclass
class FDRConfig:
    """Configuration for the FDR analysis.

    All magic numbers live here. When integrating with the GUI,
    create an FDRConfig from user inputs and pass it to FDRTool.
    """

    fy: int = 500
    """Steel yield strength (MPa)."""

    fck: int = 30
    """Default concrete grade (MPa). Used when pier-specific fck = 0."""

    text_height: int = 100
    """Text height for CAD exports."""

    cad_offset_1: int = 300
    """Gap from the pier edge to the text stack (mm)."""

    cad_offset_2: int = 200
    """Gap from the pier label up to the Required Pt% value (mm)."""

    cad_offset_3: int = 200
    """Gap from the Required Pt% value up to the As Required value (mm)."""

    wind_keywords: list[str] = field(
        default_factory=lambda: ['gx', 'gwx', 'wx', 'wy', 'gy', 'gwy']
    )
    """Substrings that flag a load combination as wind-related."""

    story_filter: Optional[str] = None
    """Set to a story name string to filter, or None for all stories."""

    show_percent_symbol: bool = True
    """Whether Required Pt% text includes a trailing '%'."""

    plot_minimum_governed_values: bool = True
    """Whether to plot Required Pt%/As values for piers governed by the
    code minimum (As_min). If False, those piers are left blank on the
    CAD/DXF/Excel outputs -- the absence of a value is the drafting
    convention for "minimum governs"."""


# %%
# =================================================================
# Pure helper functions (no state, no side-effects)
# =================================================================

def parse_namelist(ret):
    """Parse GetNameList return tuple. Returns (success:bool, count:int, names:list)."""
    if ret is None or len(ret) < 2:
        return False, 0, []
    if len(ret) >= 3 and ret[0] == 0 and isinstance(ret[1], int):
        names = list(ret[2]) if ret[2] else []
        return True, ret[1], names
    if isinstance(ret[0], int) and ret[0] > 0:
        names = list(ret[1]) if ret[1] else []
        return True, ret[0], names
    if len(ret) >= 3 and isinstance(ret[0], int) and ret[0] != 0:
        return False, 0, []
    if ret[0] == 0:
        return True, 0, []
    return False, 0, []

def is_wind_combo(combo: str, wind_keywords: list[str]) -> bool:
    """Check if a load combination name contains any wind keyword."""
    if not combo:
        return False
    lower = combo.lower()
    return any(kw in lower for kw in wind_keywords)


def natural_sort_key(s: str):
    """Generate a sort key for natural (human-friendly) sorting.
    e.g. P2 < P10, Story 1 < Story 10."""
    parts = re.split(r'(\d+)', s)
    return [int(p) if p.isdigit() else p.lower() for p in parts]


def compute_pier_corners(x1, y1, x2, y2, b):
    """Compute 4 corners of a pier rectangle from centerline midpoints + thickness.
    Returns list of (x, y) tuples, or None if degenerate."""
    dx = x2 - x1
    dy = y2 - y1
    length = math.sqrt(dx * dx + dy * dy)
    if length == 0 or b == 0:
        return None
    nx = (-dy / length) * (b / 2)
    ny = (dx / length) * (b / 2)
    return [
        (x1 + nx, y1 + ny),
        (x2 + nx, y2 + ny),
        (x2 - nx, y2 - ny),
        (x1 - nx, y1 - ny),
    ]


def boxes_overlap(a: dict, b: dict) -> bool:
    """Check if two AABB bounding boxes overlap.
    Each box is a dict with keys: xMin, xMax, yMin, yMax."""
    return (a['xMin'] < b['xMax'] and a['xMax'] > b['xMin'] and
            a['yMin'] < b['yMax'] and a['yMax'] > b['yMin'])


# -- Steel area formulas (IS 456) --

def calc_As_min(b: float, d: float) -> float:
    """Minimum steel area: 0.25% of gross area."""
    return 0.0025 * b * d


def calc_Asc(Pmin: float, fck: float, b: float, d: float, fy: float) -> float:
    """Compression steel: Asc = (|Pmin|*1000 - 0.4*fck*b*d) / (0.67*fy)."""
    if Pmin >= 0:
        return 0.0
    raw = (abs(Pmin) * 1000 - 0.4 * fck * b * d) / (0.67 * fy)
    return max(0.0, raw)


def calc_Ast(Pmax: float, fy: float) -> float:
    """Tension steel: Ast = (Pmax * 1000) / (0.87 * fy)."""
    if Pmax <= 0:
        return 0.0
    return (Pmax * 1000) / (0.87 * fy)


def governing_case(row: pd.Series) -> str:
    """Determine which steel quantity governs the design."""
    if row['As_max'] == row['Asc'] and row['Asc'] > 0:
        return 'Compression'
    elif row['As_max'] == row['Ast'] and row['Ast'] > 0:
        return 'Tension'
    else:
        return 'Minimum'


# %%
# =================================================================
# FDRTool — Main analysis engine
# =================================================================

class AnalysisRequiredError(RuntimeError):
    """Raised when the model has no analysis results available (unlocked)."""


class FDRTool:
    """FDR (Flexural Design Review) automation for ETABS pier walls.

    This class accepts a live SapModel connection (from Etabs_live.py)
    and performs the full FDR pipeline: extraction → calculation → export.

    Parameters
    ----------
    SapModel : object
        A live ETABS SapModel COM object (from Etabs_live.connect_etabs()).
    config : FDRConfig, optional
        Engineering constants. Defaults to standard IS 456 values.

    Example
    -------
    >>> from Etabs_live import connect_etabs
    >>> from fdr_tool import FDRTool, FDRConfig
    >>> _, SapModel = connect_etabs()
    >>> tool = FDRTool(SapModel, FDRConfig(fck=35, fy=500))
    >>> results = tool.run_all()
    """

    def __init__(self, SapModel, config: FDRConfig | None = None, log=None):
        self.SapModel = SapModel
        self.config = config or FDRConfig()
        self.log = log or print

        # Internal DataFrames populated by extraction/calculation steps
        self._df_raw: pd.DataFrame = pd.DataFrame()
        self._df_data: pd.DataFrame = pd.DataFrame()
        self._df_calc: pd.DataFrame = pd.DataFrame()

        # Non-fatal issues collected during extraction (e.g. piers missing
        # section properties). Surfaced by the GUI after a run.
        self.warnings: list[str] = []

        # (pier, story) -> {b, d, fck, angle, cgx, cgy}, populated by
        # extract_pier_forces() and consumed by extract_pier_coordinates().
        self._pier_geo: dict[tuple[str, str], dict] = {}

    # -----------------------------------------------------------------
    # STEP 1: Extract Pier Forces
    # -----------------------------------------------------------------

    def needs_analysis(self) -> bool:
        """Return True if the model has no analysis results available.

        The model is locked by ETABS once analysis results exist. An
        unlocked model means either analysis has never been run, or the
        model was modified since the last run.
        """
        try:
            return not self.SapModel.GetModelIsLocked()
        except Exception:
            return False

    def run_analysis(self) -> None:
        """Save the model (if it has a file path) and run analysis for all cases."""
        SM = self.SapModel

        try:
            filepath = SM.GetModelFilename()
        except Exception:
            filepath = None

        if not filepath:
            raise RuntimeError(
                "This model has never been saved. Save it in ETABS first, "
                "then run FDR analysis again."
            )

        self.log("Saving model before running analysis...")
        ret = SM.File.Save()
        if ret != 0:
            raise RuntimeError(f"Failed to save the model before analysis (ret={ret}).")

        self.log("Running analysis - this can take a while for large models...")
        ret = SM.Analyze.RunAnalysis()
        if ret != 0:
            raise RuntimeError(f"Analyze.RunAnalysis failed (ret={ret}).")

        if self.needs_analysis():
            raise RuntimeError(
                "Analysis finished but the model is still unlocked - "
                "results may not be available."
            )

        self.log("Analysis complete.")

    def get_available_combos(self) -> list[str]:
        """Return the list of response combination names defined in the model."""
        ret = self.SapModel.RespCombo.GetNameList(0, [])
        ok, n, names = parse_namelist(ret)
        if not ok:
            raise RuntimeError(
                f"RespCombo.GetNameList failed (raw={ret}). "
                "Check Define > Load Combinations in ETABS - this model may have "
                "zero response combinations defined."
            )
        return list(names)

    def extract_pier_forces(self, selected_combos: list[str] | None = None) -> pd.DataFrame:
        """Extract pier force results from ETABS and compute Pmin/Pmax envelopes.

        Pulls all load combination results, filters to bottom-location rows,
        and groups by (Story, Pier_ID) to find overall and non-wind envelopes.

        Returns
        -------
        pd.DataFrame
            Grouped data with columns: Story, Pier_ID, b, d, fck,
            Pmin, Combo_Pmin, Pmax, Combo_Pmax, and NoWind variants.
        """
        SM = self.SapModel
        cfg = self.config

        self.log("Extracting pier data from ETABS...")
        self.log("(Make sure your model is analyzed and results are available)")

        if self.needs_analysis():
            raise AnalysisRequiredError(
                "This model has no analysis results available (it is unlocked). "
                "Run analysis in ETABS, or let the FDR tool run it for you."
            )

        # Enforce kN, mm, C units to ensure calculation formulas are valid
        try:
            SM.SetPresentUnits(5)  # 5 = kN, mm, C
            self.log("  Forced ETABS API units to kN, mm, C.")
        except Exception as e:
            self.log(f"  Warning: Could not set units to kN, mm, C. Error: {e}")

        SM.Results.Setup.DeselectAllCasesAndCombosForOutput()

        available = self.get_available_combos()
        if not available:
            raise RuntimeError(
                "No response combinations are defined in this model. "
                "Define at least one under Define > Load Combinations before running FDR."
            )

        if selected_combos is None:
            combo_names = available
            self.log(f"  No combo selection given - using all {len(combo_names)} available combos.")
        else:
            invalid = [c for c in selected_combos if c not in available]
            if invalid:
                raise ValueError(f"These combo names don't exist in the model: {invalid}")
            combo_names = selected_combos
            self.log(f"  Using {len(combo_names)} user-selected combos out of {len(available)} available.")

        for combo_name in combo_names:
            SM.Results.Setup.SetComboSelectedForOutput(combo_name, True)

        # --- Get Pier Force results (no Name/ItemType args - returns ALL piers) ---
        ret = SM.Results.PierForce(
            0, [], [], [], [],
            [], [], [], [], [], []
        )
        if ret is None or len(ret) < 11:
            raise RuntimeError(f"Results.PierForce failed (raw={ret})")
            
        if ret[0] == 0 and isinstance(ret[1], int):
            num_results   = ret[1]
            story_names   = ret[2]
            pier_names    = ret[3]
            load_cases    = ret[4]
            locations     = ret[5]
            p_values      = ret[6]
            v2_values     = ret[7]
            v3_values     = ret[8]
            t_values      = ret[9]
            m2_values     = ret[10]
            m3_values     = ret[11]
        else:
            num_results   = ret[0]
            story_names   = ret[1]
            pier_names    = ret[2]
            load_cases    = ret[3]
            locations     = ret[4]
            p_values      = ret[5]
            v2_values     = ret[6]
            v3_values     = ret[7]
            t_values      = ret[8]
            m2_values     = ret[9]
            m3_values     = ret[10]

        self.log(f"  Extracted {num_results} pier force rows.")

        # --- Get Pier Section Properties ---
        ret_piers = SM.PierLabel.GetNameList(0, [])
        ok, n, pier_label_list = parse_namelist(ret_piers)
        if not ok:
            raise RuntimeError(f"PierLabel.GetNameList failed (raw={ret_piers})")

        pier_props = {}   
        for p_name in pier_label_list:
            try:
                try:
                    ret_sec = SM.PierLabel.GetSectionProperties(p_name)
                except:
                    ret_sec = SM.PierLabel.GetSectionProperties(
                        p_name, 0, [], [], [], [], [], [], [], [], [], [], [], [], [], [], []
                    )
                if ret_sec is not None and len(ret_sec) >= 7:
                    # Field order per the ETABS API docs for
                    # PierLabel.GetSectionProperties: NumberStories, StoryName,
                    # AxisAngle, NumAreaObjs, NumLineObjs, WidthBot,
                    # ThicknessBot, WidthTop, ThicknessTop, MatProp, CGBotX,
                    # CGBotY, CGBotZ, CGTopX, CGTopY, CGTopZ, [returncode].
                    # comtypes here appends the return code last rather than
                    # first, but the leading-retcode shape is handled too in
                    # case that ever differs.
                    if ret_sec[0] == 0 and isinstance(ret_sec[1], int):
                        num_stories_for_pier = ret_sec[1]
                        story_name_arr = ret_sec[2]
                        angle_arr = ret_sec[3]
                        width_bot_arr = ret_sec[6]
                        thick_bot_arr = ret_sec[7]
                        cgx_arr = ret_sec[11]
                        cgy_arr = ret_sec[12]
                    else:
                        num_stories_for_pier = ret_sec[0]
                        story_name_arr = ret_sec[1]
                        angle_arr = ret_sec[2]
                        width_bot_arr = ret_sec[5]
                        thick_bot_arr = ret_sec[6]
                        cgx_arr = ret_sec[10]
                        cgy_arr = ret_sec[11]
                    for s_idx in range(num_stories_for_pier):
                        story_nm = str(story_name_arr[s_idx]).strip().upper()
                        pier_key = str(p_name).strip().upper()
                        pier_props[(pier_key, story_nm)] = {
                            'b': float(thick_bot_arr[s_idx]),
                            'd': float(width_bot_arr[s_idx]),
                            'fck': cfg.fck,
                            'angle': float(angle_arr[s_idx]),
                            'cgx': float(cgx_arr[s_idx]),
                            'cgy': float(cgy_arr[s_idx]),
                        }
            except Exception as e:
                msg = f"Could not get section for pier {p_name}: {e}"
                self.log(f"  Warning: {msg}")
                self.warnings.append(msg)

        # --- Build raw data DataFrame (bottom-location rows only) ---
        raw_rows = []
        for i in range(num_results):
            loc = str(locations[i]).strip() if locations[i] else ""
            if loc.lower() not in ('bottom', '0', ''):
                continue
            raw_rows.append({
                'Story': str(story_names[i]).strip(),
                'Pier_ID': str(pier_names[i]).strip(),
                'Combo': str(load_cases[i]).strip(),
                'P': float(p_values[i]),
            })

        self._df_raw = pd.DataFrame(raw_rows)
        self.log(f"  Filtered to {len(self._df_raw)} bottom-location rows.")

        # --- Group by (Story, Pier_ID) → Pmin/Pmax envelopes ---
        wind_kw = cfg.wind_keywords
        grouped_data = []

        for (story, pier_id), group in self._df_raw.groupby(['Story', 'Pier_ID']):
            s_key = str(story).strip().upper()
            p_key = str(pier_id).strip().upper()
            props = pier_props.get(
                (p_key, s_key),
                {'b': 0, 'd': 0, 'fck': cfg.fck, 'angle': 0.0, 'cgx': 0.0, 'cgy': 0.0},
            )
            sorted_g = group.sort_values('P')

            min_row = sorted_g.iloc[0]
            max_row = sorted_g.iloc[-1]

            non_wind = group[~group['Combo'].apply(
                lambda c: is_wind_combo(c, wind_kw)
            )]
            if len(non_wind) > 0:
                nw_sorted = non_wind.sort_values('P')
                min_row_nw = nw_sorted.iloc[0]
                max_row_nw = nw_sorted.iloc[-1]
            else:
                min_row_nw = min_row
                max_row_nw = max_row

            grouped_data.append({
                'Story': story,
                'Pier_ID': pier_id,
                'b': props['b'],
                'd': props['d'],
                'fck': props['fck'] if props['fck'] > 0 else cfg.fck,
                'Pmin': min_row['P'],
                'Combo_Pmin': min_row['Combo'],
                'Pmax': max_row['P'],
                'Combo_Pmax': max_row['Combo'],
                'Pmin_NoWind': min_row_nw['P'],
                'Combo_Pmin_NoWind': min_row_nw['Combo'],
                'Pmax_NoWind': max_row_nw['P'],
                'Combo_Pmax_NoWind': max_row_nw['Combo'],
                'x1': 0.0, 'y1': 0.0, 'x2': 0.0, 'y2': 0.0,
            })

        self._df_data = pd.DataFrame(grouped_data)
        self._pier_geo = pier_props  # (pier, story) -> b/d/angle/cgx/cgy, reused by extract_pier_coordinates

        # Natural sort by Pier_ID
        self._df_data['_sort_key'] = self._df_data['Pier_ID'].apply(natural_sort_key)
        self._df_data = (
            self._df_data
            .sort_values('_sort_key')
            .drop(columns='_sort_key')
            .reset_index(drop=True)
        )

        self.log(f"\nEXTRACTION COMPLETE: {len(self._df_data)} unique (Story, Pier) entries")
        self.log(f"  Stories: {self._df_data['Story'].nunique()} | "
              f"Piers: {self._df_data['Pier_ID'].nunique()}")
        self.log(self._df_data[
            ['Story', 'Pier_ID', 'b', 'd', 'fck', 'Pmin', 'Pmax']
        ].head(10).to_string(index=False))

        return self._df_data.copy()

    # -----------------------------------------------------------------
    # STEP 2: Extract Pier Coordinates
    # -----------------------------------------------------------------

    def extract_pier_coordinates(self) -> pd.DataFrame:
        """Compute each pier's centerline endpoints analytically.

        PierLabel.GetSectionProperties already returns each pier's center of
        gravity (CGBotX/CGBotY) and local axis angle per story -- the pier's
        own authoritative geometry. The centerline endpoints are just that
        CG point offset by half the design width along the axis direction.

        This used to be reverse-engineered from AreaObj/PointObj corner data
        instead, which was unreliable two ways: (1) a pier label's areas
        were matched globally rather than per story, so every story but one
        got another story's geometry; (2) even for a single story, the
        matched area's raw corner extents didn't necessarily match the
        pier's own reported design width -- e.g. one pier's corners spanned
        exactly one story height (3000mm) when its actual design width was
        833mm, because the matched area was a meshed panel, not the pier
        section itself. Deriving the endpoints from the pier's own CG+angle+
        width sidesteps both problems and needs no extra ETABS calls, since
        extract_pier_forces() already fetched this data.
        """
        self.log("Computing pier coordinates from section properties...")

        df = self._df_data
        missing = 0
        for idx, row in df.iterrows():
            key = (str(row['Pier_ID']).strip().upper(), str(row['Story']).strip().upper())
            props = self._pier_geo.get(key)
            if props is None:
                missing += 1
                continue

            angle_rad = math.radians(props['angle'])
            half_len = row['d'] / 2
            dx = half_len * math.cos(angle_rad)
            dy = half_len * math.sin(angle_rad)

            df.at[idx, 'x1'] = props['cgx'] - dx
            df.at[idx, 'y1'] = props['cgy'] - dy
            df.at[idx, 'x2'] = props['cgx'] + dx
            df.at[idx, 'y2'] = props['cgy'] + dy

        coords_found = ((df['x1'] != 0) | (df['y1'] != 0)).sum()
        self.log(f"Coordinates computed for {coords_found}/{len(df)} piers.")
        if missing:
            self.log(f"  {missing} pier(s) had no matching section properties "
                  "and will be skipped in CAD export.")

        return df.copy()

    # -----------------------------------------------------------------
    # STEP 3: Run All Structural Calculations
    # -----------------------------------------------------------------

    def run_calculations(self) -> pd.DataFrame:
        """Execute all structural calculations on extracted data.

        Performs:
          1. Required Pt% (steel area) calculation
          2. 0.4 fck Capacity/Demand check (non-wind)
          3. 0.2 fck Boundary element check (non-wind)
          4. CAD label placement with collision avoidance

        Returns
        -------
        pd.DataFrame
            Full results DataFrame with all calculated columns.
        """
        cfg = self.config
        self.log("Running structural calculations...")

        # Apply story filter if set
        if cfg.story_filter:
            self._df_calc = self._df_data[
                self._df_data['Story'] == cfg.story_filter
            ].copy()
            self.log(f"  Filtered to story: {cfg.story_filter} "
                  f"({len(self._df_calc)} piers)")
        else:
            self._df_calc = self._df_data.copy()

        df = self._df_calc  # shorthand

        # ── 1. STEEL AREA CALCULATIONS ──────────────────────────────────

        df['As_min'] = df.apply(
            lambda r: calc_As_min(r['b'], r['d']), axis=1
        )
        df['Asc'] = df.apply(
            lambda r: calc_Asc(r['Pmin'], r['fck'], r['b'], r['d'], cfg.fy),
            axis=1,
        )
        df['Ast'] = df.apply(
            lambda r: calc_Ast(r['Pmax'], cfg.fy), axis=1
        )
        df['As_max'] = df[['As_min', 'Asc', 'Ast']].max(axis=1)
        df['Pt_percent'] = (df['As_max'] * 100) / (df['b'] * df['d'])
        df['Governing'] = df.apply(governing_case, axis=1)

        # ── 2. 0.4 fck C/D CHECK (Non-Wind Demands) ────────────────────

        df['Pu_capacity_04'] = (0.4 * df['fck'] * df['b'] * df['d']) / 1000
        df['Demand_04'] = df['Pmin_NoWind'].abs()
        df['CD_Ratio_04'] = df.apply(
            lambda r: (r['Pu_capacity_04'] / r['Demand_04']
                       if r['Demand_04'] > 0 else 0),
            axis=1,
        )
        df['Status_04'] = df['CD_Ratio_04'].apply(
            lambda x: 'Adequate' if x >= 1 else 'Inadequate'
        )

        # ── 3. 0.2 fck BOUNDARY CHECK (Non-Wind Demands) ───────────────

        df['Pu_02fck'] = (0.2 * df['fck'] * df['b'] * df['d']) / 1000
        df['Demand_02'] = df['Pmin_NoWind'].abs()
        df['CD_Ratio_02'] = df.apply(
            lambda r: (r['Pu_02fck'] / r['Demand_02']
                       if r['Demand_02'] > 0 else 0),
            axis=1,
        )
        df['Ductile_Detailing_Reqd'] = df['CD_Ratio_02'].apply(
            lambda x: 'Yes' if x < 1 else 'No'
        )

        # ── 4. CAD LABEL PLACEMENT (collision avoidance) ────────────────

        self._place_cad_labels(df)

        # ── Summary ─────────────────────────────────────────────────────

        summary = self.get_summary()
        self.log(f"\n{'='*55}")
        self.log(f"  CALCULATION SUMMARY")
        self.log(f"{'='*55}")
        self.log(f"  Total Piers Analyzed:   {summary['total_piers']}")
        self.log(f"  Max Pt%:                {summary['max_pt']:.2f}%")
        self.log(f"  0.4fck Inadequate:      {summary['inadequate_04']}")
        self.log(f"  Ductile Detailing Reqd: {summary['ductile_reqd']}")
        self.log(f"{'='*55}")

        self._df_calc = df
        return self._df_calc.copy()

    def _place_cad_labels(self, df: pd.DataFrame) -> None:
        """Compute label positions for CAD export using uniform spacing and collision avoidance."""
        cfg = self.config
        
        # We will use cad_offset_1 as the gap from the pier edge to the text center.
        gap = cfg.cad_offset_1
        
        # Calculate bounding boxes of piers to avoid overlaps
        pier_boxes = []
        for _, row in df.iterrows():
            corners = compute_pier_corners(
                row['x1'], row['y1'], row['x2'], row['y2'], row['b']
            )
            if corners is None:
                continue
            xs = [c[0] for c in corners]
            ys = [c[1] for c in corners]
            pier_boxes.append({
                'id': f"{row['Story']}_{row['Pier_ID']}",
                'xMin': min(xs), 'xMax': max(xs),
                'yMin': min(ys), 'yMax': max(ys),
            })
            
        placed_labels = []
        
        CHAR_W = 70
        TEXT_H = cfg.text_height

        for idx, row in df.iterrows():
            x1, y1, x2, y2, b = row['x1'], row['y1'], row['x2'], row['y2'], row['b']
            mid_x = (x1 + x2) / 2
            mid_y = (y1 + y2) / 2
            dx = abs(x2 - x1)
            dy = abs(y2 - y1)
            
            pier_key = f"{row['Story']}_{row['Pier_ID']}"
            other_boxes = [box for box in pier_boxes if box['id'] != pier_key]
            
            is_horiz = (dy < dx) or (y1 == y2)
            
            # The edge of the pier from the center:
            edge_x = max(dx / 2, b / 2) if not is_horiz else b / 2
            edge_y = max(dy / 2, b / 2) if is_horiz else b / 2
            
            # Try to place text block 
            if is_horiz:
                # Primary: Above -> Below -> Right -> Left
                candidates = [
                    (mid_x, mid_y + edge_y + gap),
                    (mid_x, mid_y - edge_y - gap),
                    (mid_x + edge_x + gap, mid_y),
                    (mid_x - edge_x - gap, mid_y),
                ]
            else:
                # Primary: Right -> Left -> Above -> Below
                candidates = [
                    (mid_x + edge_x + gap, mid_y),
                    (mid_x - edge_x - gap, mid_y),
                    (mid_x, mid_y + edge_y + gap),
                    (mid_x, mid_y - edge_y - gap),
                ]
                
            max_text_len = max(len(str(row['Pier_ID'])), 6)
            text_half_w = (max_text_len * CHAR_W) / 2
            
            def make_label_box(lx, ly):
                # Stack grows upward from the pier label: PierLabel at ly,
                # RequiredPt above it, RequiredAs above that.
                top_y = ly + cfg.cad_offset_2 + cfg.cad_offset_3
                return {
                    'xMin': lx - text_half_w,
                    'xMax': lx + text_half_w,
                    'yMin': min(ly, top_y) - TEXT_H / 2,
                    'yMax': max(ly, top_y) + TEXT_H / 2,
                }
                
            label_x, label_y = candidates[0]
            placed = False
            
            for cx, cy in candidates:
                text_box = make_label_box(cx, cy)
                hits_pier = any(boxes_overlap(text_box, pbox) for pbox in other_boxes)
                hits_label = any(boxes_overlap(text_box, lbox) for lbox in placed_labels)
                if not hits_pier and not hits_label:
                    label_x, label_y = cx, cy
                    placed_labels.append(text_box)
                    placed = True
                    break
                    
            if not placed:
                placed_labels.append(make_label_box(label_x, label_y))
                
            df.at[idx, 'labelX'] = label_x
            df.at[idx, 'labelY'] = label_y
            df.at[idx, 'valueX'] = label_x
            df.at[idx, 'valueY'] = label_y + cfg.cad_offset_2
            df.at[idx, 'asValueX'] = label_x
            df.at[idx, 'asValueY'] = label_y + cfg.cad_offset_2 + cfg.cad_offset_3

    # -----------------------------------------------------------------
    # Summary
    # -----------------------------------------------------------------

    def get_summary(self) -> dict:
        """Return summary statistics from the calculation results.

        Returns
        -------
        dict
            Keys: total_piers, stories, max_pt, inadequate_04, ductile_reqd,
            tension_count, compression_count, minimum_count, no_compression.
        """
        df = self._df_calc
        return {
            'total_piers': len(df),
            'stories': int(df['Story'].nunique()) if len(df) > 0 else 0,
            'max_pt': df['Pt_percent'].max() if len(df) > 0 else 0.0,
            'inadequate_04': int((df['Status_04'] == 'Inadequate').sum()),
            'ductile_reqd': int(
                (df['Ductile_Detailing_Reqd'] == 'Yes').sum()
            ),
            'tension_count': int((df['Governing'] == 'Tension').sum()),
            'compression_count': int((df['Governing'] == 'Compression').sum()),
            'minimum_count': int((df['Governing'] == 'Minimum').sum()),
            'no_compression': int((df['Pmin'] >= 0).sum()) if len(df) > 0 else 0,
        }

    # -----------------------------------------------------------------
    # Print detailed results (matches original Cells 7-9)
    # -----------------------------------------------------------------

    def print_pt_results(self) -> None:
        """Print the Required Pt% analysis table (Cell 7)."""
        df = self._df_calc
        pt_cols = [
            'Story', 'Pier_ID', 'b', 'd', 'fck',
            'Pmin', 'Combo_Pmin', 'Pmax', 'Combo_Pmax',
            'As_min', 'Asc', 'Ast', 'As_max', 'Pt_percent', 'Governing',
        ]

        self.log("=" * 100)
        self.log("  OBJECTIVE 1: REQUIRED Pt% ANALYSIS")
        self.log("  Asc = compression steel (from Pmin) | "
              "Ast = tension steel (from Pmax)")
        self.log(f"  As_min = 0.0025 x b x d | As_max = max(Asc, Ast, As_min) "
              f"| fy = {self.config.fy} MPa")
        self.log("=" * 100)

        display = df[pt_cols].copy()
        display['Pt_percent'] = display['Pt_percent'].round(3)
        for col in ('As_min', 'Asc', 'Ast', 'As_max'):
            display[col] = display[col].round(0)
        for col in ('Pmin', 'Pmax'):
            display[col] = display[col].round(1)
        self.log(display.to_string(index=False))

        tension = df[df['Governing'] == 'Tension']
        compression = df[df['Governing'] == 'Compression']
        minimum = df[df['Governing'] == 'Minimum']
        self.log(f"\n  Governed by Tension:     {len(tension)}")
        self.log(f"  Governed by Compression: {len(compression)}")
        self.log(f"  Governed by Minimum:     {len(minimum)}")
        if len(tension) > 0:
            self.log(f"  Max Pt% (Tension):       {tension['Pt_percent'].max():.2f}%")
        if len(compression) > 0:
            self.log(f"  Max Pt% (Compression):   "
                  f"({compression['Pt_percent'].max():.2f}%)")

    def print_04fck_results(self) -> None:
        """Print the 0.4 fck C/D check table (Cell 8)."""
        df = self._df_calc
        cd04_cols = [
            'Story', 'Pier_ID', 'Pu_capacity_04', 'Demand_04',
            'Combo_Pmin_NoWind', 'CD_Ratio_04', 'Status_04',
        ]

        self.log("=" * 100)
        self.log("  OBJECTIVE 2: 0.4 fck CAPACITY/DEMAND CHECK")
        self.log("  Pu_capacity = 0.4 x fck x b x d / 1000  |  "
              "Demand = |Pmin_NoWind|")
        self.log("  Status: Adequate if C/D >= 1.0")
        self.log("=" * 100)

        display = df[cd04_cols].copy()
        display['Pu_capacity_04'] = display['Pu_capacity_04'].round(1)
        display['Demand_04'] = display['Demand_04'].round(1)
        display['CD_Ratio_04'] = display['CD_Ratio_04'].round(2)
        self.log(display.to_string(index=False))

        inadequate = (df['Status_04'] == 'Inadequate').sum()
        self.log(f"\n  ADEQUATE:   {(df['Status_04'] == 'Adequate').sum()}")
        self.log(f"  INADEQUATE: {inadequate}")

        if inadequate > 0:
            self.log("\n  INADEQUATE PIERS:")
            inad = df[df['Status_04'] == 'Inadequate'][
                ['Story', 'Pier_ID', 'CD_Ratio_04', 'Demand_04',
                 'Pu_capacity_04']
            ]
            self.log(inad.to_string(index=False))

    def print_02fck_results(self) -> None:
        """Print the 0.2 fck Boundary check table (Cell 9)."""
        df = self._df_calc
        cd02_cols = [
            'Story', 'Pier_ID', 'Demand_02', 'Combo_Pmin_NoWind',
            'Pu_02fck', 'CD_Ratio_02', 'Ductile_Detailing_Reqd',
        ]

        self.log("=" * 100)
        self.log("  OBJECTIVE 3: 0.2 fck BOUNDARY CHECK (DUCTILE DETAILING)")
        self.log("  Pu_02fck = 0.2 x fck x b x d / 1000  |  "
              "Demand = |Pmin_NoWind|")
        self.log("  Ductile detailing required if C/D < 1.0")
        self.log("=" * 100)

        display = df[cd02_cols].copy()
        display['Demand_02'] = display['Demand_02'].round(1)
        display['Pu_02fck'] = display['Pu_02fck'].round(1)
        display['CD_Ratio_02'] = display['CD_Ratio_02'].round(2)
        self.log(display.to_string(index=False))

        ductile_reqd = (df['Ductile_Detailing_Reqd'] == 'Yes').sum()
        self.log(f"\n  Ductile Detailing NOT Required: "
              f"{(df['Ductile_Detailing_Reqd'] == 'No').sum()}")
        self.log(f"  Ductile Detailing REQUIRED:     {ductile_reqd}")

        if ductile_reqd > 0:
            self.log("\n  PIERS REQUIRING DUCTILE DETAILING:")
            duct = df[df['Ductile_Detailing_Reqd'] == 'Yes'][
                ['Story', 'Pier_ID', 'CD_Ratio_02', 'Demand_02', 'Pu_02fck']
            ]
            self.log(duct.to_string(index=False))

    # -----------------------------------------------------------------
    # Shared label text formatting (used by export_cad/export_excel/export_dxf)
    # -----------------------------------------------------------------

    def _pt_label(self, row) -> str | None:
        """Formatted Required Pt% text for one row, or None to plot nothing."""
        cfg = self.config
        if row['Governing'] == 'Minimum' and not cfg.plot_minimum_governed_values:
            return None
        is_comp = row['As_max'] == row['Asc'] and row['Asc'] > 0
        suffix = '%' if cfg.show_percent_symbol else ''
        value = f"{row['Pt_percent']:.2f}{suffix}"
        return f"({value})" if is_comp else value

    def _as_label(self, row) -> str | None:
        """Formatted As Required text for one row, or None to plot nothing."""
        cfg = self.config
        if row['Governing'] == 'Minimum' and not cfg.plot_minimum_governed_values:
            return None
        is_comp = row['As_max'] == row['Asc'] and row['Asc'] > 0
        value = f"{row['As_max']:.0f}"
        return f"({value})" if is_comp else value

    # -----------------------------------------------------------------
    # EXPORT: CAD Scripts
    # -----------------------------------------------------------------

    def export_cad(self, output_dir: str | None = None) -> list[str]:
        """Generate AutoCAD .scr script files for all result categories.

        Parameters
        ----------
        output_dir : str, optional
            Directory to save .scr files. Defaults to ./cad_export.

        Returns
        -------
        list[str]
            Paths to the generated .scr files.
        """
        if output_dir is None:
            output_dir = os.path.join(
                os.path.dirname(os.path.abspath(__file__)), "cad_export"
            )
        os.makedirs(output_dir, exist_ok=True)
        timestamp = datetime.now().strftime("%Y%m%d_%H%M")
        df = self._df_calc
        generated_files = []
        TEXT_H = self.config.text_height

        # Creates (or redefines) a "Calibri" text style and makes it current,
        # so the plain "-text" lines below -- which don't specify a style --
        # pick it up instead of inheriting whatever style happens to already
        # be current in the drawing they're pasted into. Field order follows
        # the "-STYLE" command prompts: name, font, height, width factor,
        # obliquing angle, backwards?, upside-down?, vertical? (blank = keep
        # that prompt's default).
        style_setup = "-style\t" + "\t".join(
            ["Calibri", "calibri.ttf", "", "", "", "", "", ""]
        )

        def _save(name, rows, set_style=False):
            filename = f"FDR_{name}_{timestamp}.scr"
            filepath = os.path.join(output_dir, filename)
            out_rows = ([style_setup] + rows) if set_style else rows
            with open(filepath, 'w') as f:
                f.write('\n'.join(out_rows))
            self.log(f"  {name}: {len(rows)} rows -> {filename}")
            generated_files.append(filepath)
            return filepath

        # --- Pier Labels ---
        label_rows = []
        for _, r in df.iterrows():
            if 'labelX' not in r or pd.isna(r.get('labelX', None)):
                continue
            label_rows.append(
                f"-text\t{r['labelX']:.1f},{r['labelY']:.1f}\t{TEXT_H}\t0\t"
                f"{r['Pier_ID']}"
            )

        # --- Required Pt% ---
        pt_rows = []
        for _, r in df.iterrows():
            if 'valueX' not in r or pd.isna(r.get('valueX', None)):
                continue
            pt_label = self._pt_label(r)
            if pt_label is None:
                continue
            pt_rows.append(
                f"-text\t{r['valueX']:.1f},{r['valueY']:.1f}\t{TEXT_H}\t0\t"
                f"{pt_label}"
            )

        # --- 0.4 fck C/D ---
        cd04_rows = []
        for _, r in df.iterrows():
            if 'labelX' not in r or pd.isna(r.get('labelX', None)):
                continue
            cd04_rows.append(
                f"-text\t{r['labelX']:.1f},{r['labelY']:.1f}\t{TEXT_H}\t0\t"
                f"{r['CD_Ratio_04']:.2f}"
            )

        # --- 0.2 fck C/D ---
        cd02_rows = []
        for _, r in df.iterrows():
            if 'labelX' not in r or pd.isna(r.get('labelX', None)):
                continue
            cd02_rows.append(
                f"-text\t{r['labelX']:.1f},{r['labelY']:.1f}\t{TEXT_H}\t0\t"
                f"{r['CD_Ratio_02']:.2f}"
            )

        # --- As Required ---
        as_rows = []
        for _, r in df.iterrows():
            if 'asValueX' not in r or pd.isna(r.get('asValueX', None)):
                continue
            as_label = self._as_label(r)
            if as_label is None:
                continue
            as_rows.append(
                f"-text\t{r['asValueX']:.1f},{r['asValueY']:.1f}\t{TEXT_H}\t0\t"
                f"{as_label}"
            )

        # --- Pier Rectangles ---
        rect_rows = []
        skipped_piers = []
        for _, r in df.iterrows():
            x1, y1, x2, y2, b = r['x1'], r['y1'], r['x2'], r['y2'], r['b']
            if x1 == 0 and y1 == 0 and x2 == 0 and y2 == 0:
                skipped_piers.append(r['Pier_ID'])
                continue
            if b == 0:
                skipped_piers.append(r['Pier_ID'])
                continue
            corners = compute_pier_corners(x1, y1, x2, y2, b)
            if corners is None:
                skipped_piers.append(r['Pier_ID'])
                continue
            c1, c2, c3, c4 = corners
            rect_rows.append(
                f"PLINE {c1[0]:.1f},{c1[1]:.1f} {c2[0]:.1f},{c2[1]:.1f} "
                f"{c3[0]:.1f},{c3[1]:.1f} {c4[0]:.1f},{c4[1]:.1f} C"
            )

        self.log(f"\n{'='*60}")
        self.log(f"  CAD EXPORT -- saved to: {output_dir}")
        self.log(f"{'='*60}")

        _save("PierLabels", label_rows, set_style=True)
        _save("RequiredPt", pt_rows, set_style=True)
        _save("CD_04fck", cd04_rows, set_style=True)
        _save("CD_02fck", cd02_rows, set_style=True)
        _save("As_Required", as_rows, set_style=True)
        _save("PierRectangles", rect_rows)

        if skipped_piers:
            self.log(f"\n  {len(skipped_piers)} pier(s) skipped for rectangles "
                  f"(no coords/thickness)")

        self.log(f"\nAll CAD scripts saved. "
              f"Paste contents into AutoCAD command line.")

        return generated_files

    # -----------------------------------------------------------------
    # EXPORT: Excel Report
    # -----------------------------------------------------------------

    def export_excel(self, output_path: str | None = None) -> str:
        """Export full analysis results to a multi-sheet Excel workbook.

        Parameters
        ----------
        output_path : str, optional
            Full path for the .xlsx file. If None, prompts via filedialog.

        Returns
        -------
        str
            Path to the generated Excel file.
        """
        try:
            import openpyxl
            from openpyxl.styles import PatternFill, Border, Side, Font
        except ImportError:
            self.log("Installing openpyxl for Excel export...")
            import subprocess
            subprocess.check_call(['pip', 'install', 'openpyxl'])
            import openpyxl
            from openpyxl.styles import PatternFill, Border, Side, Font

        if output_path is None:
            import tkinter as tk
            from tkinter import filedialog
            
            root = tk.Tk()
            root.withdraw()
            
            try:
                model_dir = os.path.dirname(self.SapModel.GetModelFilename())
            except Exception:
                model_dir = os.path.expanduser("~")
                
            output_path = filedialog.asksaveasfilename(
                title="Save FDR Excel Report",
                initialdir=model_dir,
                initialfile=self.default_excel_name(),
                defaultextension=".xlsx",
                filetypes=[("Excel workbook", "*.xlsx"), ("All files", "*.*")]
            )
            root.destroy()
            if not output_path:
                self.log("Excel export cancelled.")
                return ""

        df = self._df_calc
        export_cols = [
            'Story', 'Pier_ID', 'b', 'd', 'fck',
            'x1', 'y1', 'x2', 'y2',
            'Pmin', 'Combo_Pmin', 'Pmax', 'Combo_Pmax',
            'Pmin_NoWind', 'Combo_Pmin_NoWind',
            'Pmax_NoWind', 'Combo_Pmax_NoWind',
            'As_min', 'Asc', 'Ast', 'As_max', 'Pt_percent', 'Governing',
            'Pu_capacity_04', 'Demand_04', 'CD_Ratio_04', 'Status_04',
            'Pu_02fck', 'Demand_02', 'CD_Ratio_02', 'Ductile_Detailing_Reqd',
        ]
        export_cols = [c for c in export_cols if c in df.columns]

        # Prepare separate AutoCAD Commands dataframes
        cad_rect_rows = []
        cad_label_rows = []
        cad_pt_rows = []
        cad_as_rows = []
        cad_04_rows = []
        cad_02_rows = []
        
        TEXT_H = self.config.text_height

        if not df.empty:
            for _, r in df.iterrows():
                # Rectangle
                x1, y1, x2, y2, b = r['x1'], r['y1'], r['x2'], r['y2'], r['b']
                if not (x1 == 0 and y1 == 0 and x2 == 0 and y2 == 0) and b > 0:
                    corners = compute_pier_corners(x1, y1, x2, y2, b)
                    if corners is not None:
                        c1, c2, c3, c4 = corners
                        cad_rect_rows.append({
                            'Pier_ID': r['Pier_ID'], 'Command': 'PLINE',
                            'Point1': f"{c1[0]:.1f},{c1[1]:.1f}",
                            'Point2': f"{c2[0]:.1f},{c2[1]:.1f}",
                            'Point3': f"{c3[0]:.1f},{c3[1]:.1f}",
                            'Point4': f"{c4[0]:.1f},{c4[1]:.1f}",
                            'Close': 'C',
                        })

                # Label
                if 'labelX' in r and not pd.isna(r['labelX']):
                    label_xy = f"{r['labelX']:.1f},{r['labelY']:.1f}"
                    cad_label_rows.append({
                        'Pier_ID': r['Pier_ID'], 'Command': '-text',
                        'Point': label_xy,
                        'Height': TEXT_H, 'Rotation': 0, 'Text': r['Pier_ID'],
                    })

                    # 0.4fck and 0.2fck use the base label coordinate because they go on separate drawings
                    cad_04_rows.append({
                        'Pier_ID': r['Pier_ID'], 'Command': '-text',
                        'Point': label_xy,
                        'Height': TEXT_H, 'Rotation': 0,
                        'Text': round(r['CD_Ratio_04'], 2),
                    })
                    cad_02_rows.append({
                        'Pier_ID': r['Pier_ID'], 'Command': '-text',
                        'Point': label_xy,
                        'Height': TEXT_H, 'Rotation': 0,
                        'Text': round(r['CD_Ratio_02'], 2),
                    })

                # Required Pt%
                if 'valueX' in r and not pd.isna(r['valueX']):
                    pt_label = self._pt_label(r)
                    if pt_label is not None:
                        cad_pt_rows.append({
                            'Pier_ID': r['Pier_ID'], 'Command': '-text',
                            'Point': f"{r['valueX']:.1f},{r['valueY']:.1f}",
                            'Height': TEXT_H, 'Rotation': 0, 'Text': pt_label,
                        })

                # As Required
                if 'asValueX' in r and not pd.isna(r['asValueX']):
                    as_label = self._as_label(r)
                    if as_label is not None:
                        cad_as_rows.append({
                            'Pier_ID': r['Pier_ID'], 'Command': '-text',
                            'Point': f"{r['asValueX']:.1f},{r['asValueY']:.1f}",
                            'Height': TEXT_H, 'Rotation': 0, 'Text': as_label,
                        })

        df_cad_rect = pd.DataFrame(cad_rect_rows)
        df_cad_label = pd.DataFrame(cad_label_rows)
        df_cad_pt = pd.DataFrame(cad_pt_rows)
        df_cad_as = pd.DataFrame(cad_as_rows)
        df_cad_04 = pd.DataFrame(cad_04_rows)
        df_cad_02 = pd.DataFrame(cad_02_rows)

        # Mapping for unit-aware headers
        rename_map = {
            'b': 'b (mm)',
            'd': 'd (mm)',
            'fck': 'fck (MPa)',
            'x1': 'x1 (mm)', 'y1': 'y1 (mm)', 'x2': 'x2 (mm)', 'y2': 'y2 (mm)',
            'Pmin': 'Pmin (kN)', 'Pmax': 'Pmax (kN)',
            'Pmin_NoWind': 'Pmin_NoWind (kN)', 'Pmax_NoWind': 'Pmax_NoWind (kN)',
            'As_min': 'As_min (mm2)', 'Asc': 'Asc (mm2)', 'Ast': 'Ast (mm2)', 'As_max': 'As_max (mm2)',
            'Pu_capacity_04': 'Pu_capacity_04 (kN)', 'Demand_04': 'Demand_04 (kN)',
            'Pu_02fck': 'Pu_02fck (kN)', 'Demand_02': 'Demand_02 (kN)',
        }
        
        df_export = df.rename(columns=rename_map)
        export_cols_full = [rename_map.get(c, c) for c in export_cols if c in df.columns]

        with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
            df_export[export_cols_full].to_excel(writer, sheet_name='Full Analysis', index=False)
            
            pt_cols = [rename_map.get(c, c) for c in ['Story', 'Pier_ID', 'b', 'd', 'fck', 'Pmin', 'Combo_Pmin', 'Pmax', 'Combo_Pmax', 'As_min', 'Asc', 'Ast', 'As_max', 'Pt_percent', 'Governing']]
            df_export[pt_cols].to_excel(writer, sheet_name='Required Pt', index=False)
            
            cd04_cols = [rename_map.get(c, c) for c in ['Story', 'Pier_ID', 'Pu_capacity_04', 'Demand_04', 'Combo_Pmin_NoWind', 'CD_Ratio_04', 'Status_04']]
            df_export[cd04_cols].to_excel(writer, sheet_name='0.4fck CD Check', index=False)
            
            cd02_cols = [rename_map.get(c, c) for c in ['Story', 'Pier_ID', 'Demand_02', 'Combo_Pmin_NoWind', 'Pu_02fck', 'CD_Ratio_02', 'Ductile_Detailing_Reqd']]
            df_export[cd02_cols].to_excel(writer, sheet_name='0.2fck Boundary', index=False)
            
            if not df_cad_rect.empty: df_cad_rect.to_excel(writer, sheet_name='CAD - Pier Rectangles', index=False)
            if not df_cad_label.empty: df_cad_label.to_excel(writer, sheet_name='CAD - Pier Labels', index=False)
            if not df_cad_pt.empty: df_cad_pt.to_excel(writer, sheet_name='CAD - Required Pt', index=False)
            if not df_cad_as.empty: df_cad_as.to_excel(writer, sheet_name='CAD - As Required', index=False)
            if not df_cad_04.empty: df_cad_04.to_excel(writer, sheet_name='CAD - 0.4fck CD', index=False)
            if not df_cad_02.empty: df_cad_02.to_excel(writer, sheet_name='CAD - 0.2fck CD', index=False)

            # Apply formatting. Bounded regardless of sheet size so a large
            # model can never turn this into a multi-minute, file-corrupting
            # write (this used to also embed the raw per-combo envelope here
            # -- hundreds of thousands of rows -- which is why it's now a
            # separate CSV instead, see below).
            workbook = writer.book
            thin_border = Border(left=Side(style='thin'), right=Side(style='thin'),
                                 top=Side(style='thin'), bottom=Side(style='thin'))
            header_fill = PatternFill(start_color="FCE4D6", end_color="FCE4D6", fill_type="solid")
            header_font = Font(bold=True)

            BORDER_ROW_LIMIT = 20000    # skip per-cell borders past this many rows
            WIDTH_SAMPLE_ROWS = 200     # column width only needs a sample

            for sheet_name in workbook.sheetnames:
                ws = workbook[sheet_name]
                apply_borders = ws.max_row <= BORDER_ROW_LIMIT

                for col_idx in range(1, ws.max_column + 1):
                    cell = ws.cell(row=1, column=col_idx)
                    cell.fill = header_fill
                    cell.font = header_font
                    if apply_borders:
                        cell.border = thin_border

                if apply_borders and ws.max_row > 1:
                    for row in ws.iter_rows(min_row=2, max_row=ws.max_row,
                                            min_col=1, max_col=ws.max_column):
                        for cell in row:
                            cell.border = thin_border
                else:
                    self.log(f"  ({sheet_name}: {ws.max_row} rows -- skipped "
                          f"cell borders to keep export fast)")

                # Auto-adjust column widths from a bounded sample of rows
                sample_limit = min(ws.max_row, WIDTH_SAMPLE_ROWS)
                for col in ws.iter_cols(min_row=1, max_row=sample_limit,
                                       max_col=ws.max_column):
                    max_length = 0
                    column = col[0].column_letter
                    for cell in col:
                        try:
                            if len(str(cell.value)) > max_length:
                                max_length = len(str(cell.value))
                        except Exception:
                            pass
                    ws.column_dimensions[column].width = max_length + 2

        self.log(f"Excel report saved: {output_path}")
        self.log("  Sheets: Full Analysis | Required Pt | 0.4fck CD Check | 0.2fck Boundary | CAD sheets")

        # Raw per-combo envelope (one row per pier/story/combo, can be huge
        # for a real model -- hundreds of thousands of rows). Writing this
        # into the xlsx via openpyxl used to make the file enormous and slow
        # enough to write that it could come out corrupt. A CSV handles the
        # same volume in a couple of seconds.
        if len(self._df_raw) > 0:
            raw_csv_path = os.path.splitext(output_path)[0] + "_RawEnvelope.csv"
            raw_export = self._df_raw.rename(columns={'P': 'P (kN)'})
            raw_export.to_csv(raw_csv_path, index=False)
            self.log(f"  Raw envelope ({len(raw_export)} rows) saved separately: {raw_csv_path}")

        return output_path

    def get_stories(self) -> list[str]:
        """Return a list of all story names defined in the model."""
        ret = self.SapModel.Story.GetNameList(0, [])
        ok, n, names = parse_namelist(ret)
        if not ok:
            return []
        return list(names)

    def stories_in_results(self) -> list[str]:
        """Return a list of stories present in the calculation results."""
        if self._df_calc is None or self._df_calc.empty:
            return []
        return self._df_calc['Story'].unique().tolist()

    def default_excel_name(self) -> str:
        """Default filename for Excel export."""
        timestamp = datetime.now().strftime("%Y%m%d_%H%M")
        return f"FDR_Report_{timestamp}.xlsx"

    def default_dxf_name(self, story: str) -> str:
        """Default filename for DXF export."""
        timestamp = datetime.now().strftime("%Y%m%d_%H%M")
        s_safe = "".join(c if c.isalnum() else "_" for c in story)
        return f"FDR_{s_safe}_{timestamp}.dxf"

    def export_dxf(self, filepath: str, story: str) -> str:
        """Export analysis results for a specific story to DXF using ezdxf.

        Parameters
        ----------
        filepath : str
            Full path for the .dxf file.
        story : str
            The name of the story to export.

        Returns
        -------
        str
            Path to the generated DXF file.
        """
        try:
            import ezdxf
        except ImportError:
            self.log("Installing ezdxf for DXF export...")
            import subprocess
            subprocess.check_call(['pip', 'install', 'ezdxf'])
            import ezdxf

        df = self._df_calc[self._df_calc['Story'] == story]
        if df.empty:
            self.log(f"No results found for story: {story}")
            return ""

        doc = ezdxf.new('R2010')
        doc.header['$INSUNITS'] = 4  # millimeters
        msp = doc.modelspace()

        # Add layers
        doc.layers.add(name="PierRectangles", color=ezdxf.colors.WHITE)
        doc.layers.add(name="PierLabels", color=ezdxf.colors.YELLOW)
        doc.layers.add(name="RequiredPt", color=ezdxf.colors.CYAN)
        doc.layers.add(name="AsRequired", color=ezdxf.colors.MAGENTA)

        doc.styles.add("Calibri", font="calibri.ttf")

        TEXT_H = self.config.text_height

        for _, r in df.iterrows():
            # Pier Rectangle
            x1, y1, x2, y2, b = r['x1'], r['y1'], r['x2'], r['y2'], r['b']
            if not (x1 == 0 and y1 == 0 and x2 == 0 and y2 == 0) and b > 0:
                corners = compute_pier_corners(x1, y1, x2, y2, b)
                if corners is not None:
                    c1, c2, c3, c4 = corners
                    msp.add_lwpolyline([c1, c2, c3, c4], close=True, dxfattribs={'layer': 'PierRectangles'})

            # Label
            if 'labelX' in r and not pd.isna(r['labelX']):
                msp.add_text(str(r['Pier_ID']), dxfattribs={'layer': 'PierLabels', 'height': TEXT_H, 'style': 'Calibri'}).set_placement(
                    (r['labelX'], r['labelY']), align=ezdxf.enums.TextEntityAlignment.CENTER
                )

            # Required Pt%
            if 'valueX' in r and not pd.isna(r['valueX']):
                pt_label = self._pt_label(r)
                if pt_label is not None:
                    msp.add_text(pt_label, dxfattribs={'layer': 'RequiredPt', 'height': TEXT_H, 'style': 'Calibri'}).set_placement(
                        (r['valueX'], r['valueY']), align=ezdxf.enums.TextEntityAlignment.CENTER
                    )

            # As Required
            if 'asValueX' in r and not pd.isna(r['asValueX']):
                as_label = self._as_label(r)
                if as_label is not None:
                    msp.add_text(as_label, dxfattribs={'layer': 'AsRequired', 'height': TEXT_H, 'style': 'Calibri'}).set_placement(
                        (r['asValueX'], r['asValueY']), align=ezdxf.enums.TextEntityAlignment.CENTER
                    )

        doc.saveas(filepath)
        self.log(f"DXF saved: {filepath}")
        return filepath

    # -----------------------------------------------------------------
    # RUN ALL — One-shot full pipeline
    # -----------------------------------------------------------------

    def run_all(self, selected_combos: list[str] | None = None, output_dir: str | None = None) -> dict:
        """Execute the complete FDR pipeline in one call.

        Sequence: extract forces → extract coords → calculate →
        print results → export CAD → export Excel.

        Parameters
        ----------
        selected_combos : list[str], optional
            Subset of combinations to extract, or None for all.
        output_dir : str, optional
            Base directory for exports. Defaults to script directory.

        Returns
        -------
        dict
            Keys: 'summary', 'df_results', 'cad_files', 'excel_path'.
        """
        # Step 1: Extract
        self.extract_pier_forces(selected_combos)

        # Step 2: Coordinates
        self.extract_pier_coordinates()

        # Step 3: Calculate
        self.run_calculations()

        # Step 4: Print detailed results
        self.print_pt_results()
        self.print_04fck_results()
        self.print_02fck_results()

        # Step 5: Export
        cad_files = self.export_cad(
            os.path.join(output_dir, "cad_export") if output_dir else None
        )
        excel_path = self.export_excel(
            os.path.join(output_dir, f"FDR_Report_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx")
            if output_dir else None
        )

        return {
            'summary': self.get_summary(),
            'df_results': self._df_calc.copy(),
            'cad_files': cad_files,
            'excel_path': excel_path,
        }


# %%
# =================================================================
# CELL 1: CONNECT TO ETABS (Standalone / Interactive)
# =================================================================
if __name__ == '__main__':
    import comtypes.client
    import psutil

    # 1. Find the running ETABS process
    etabs_processes = [
        p.info for p in psutil.process_iter(['pid', 'name'])
        if p.info['name'] and 'ETABS' in p.info['name'].upper()
    ]

    if not etabs_processes:
        raise RuntimeError("No running ETABS.exe process found.")
    elif len(etabs_processes) > 1:
        print("Multiple ETABS processes found:")
        for p in etabs_processes:
            print(f"  PID {p['pid']}  ({p['name']})")
        raise RuntimeError("Multiple ETABS instances running — specify the correct PID manually instead of auto-picking.")
    else:
        pid = etabs_processes[0]['pid']
        print(f"Found ETABS process: PID {pid}")

    # 2. Attach to that exact process
    helper = comtypes.client.CreateObject('ETABSv1.Helper')
    helper = helper.QueryInterface(comtypes.gen.ETABSv1.cHelper)

    myETABSObject = helper.GetObjectProcess("CSI.ETABS.API.ETABSObject", pid)

    if myETABSObject is None:
        raise RuntimeError(f"GetObjectProcess returned None for PID {pid} — attach failed.")

    myETABSObject = myETABSObject.QueryInterface(comtypes.gen.ETABSv1.cOAPI)
    SapModel = myETABSObject.SapModel

    print("Attached successfully. Active file:", SapModel.GetModelFilename())

    # =================================================================
    # CELL 1.5: CHECK AVAILABLE COMBOS
    # =================================================================
    config = FDRConfig(
        fy=500,
        fck=30,
        story_filter=None,  # Set to e.g. "Story 1" to filter
    )
    tool = FDRTool(SapModel, config)

    available = tool.get_available_combos()
    print(f"{len(available)} combos found:")
    for c in available:
        print(" ", c)

    # Example: If you only want to run specific combos, uncomment and list them:
    # my_combos = ["STD_G+40_ULS_D1a", "STD_G+40_ULS_D1b"]
    my_combos = None

    # =================================================================
    # CELL 2: RUN FDR ANALYSIS
    # =================================================================
    print("=" * 60)
    print("  FDR AUTOMATION TOOL — Running Analysis")
    print("=" * 60)

    results = tool.run_all(selected_combos=my_combos)

    print("\n" + "=" * 60)
    print("  ALL DONE!")
    print("=" * 60)
